from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from PIL import Image
import json

ROOT=Path(__file__).resolve().parent
DOCS=ROOT.parent
prs=Presentation()
prs.slide_width=Inches(13.333333)
prs.slide_height=Inches(7.5)
prs.core_properties.title="Computação Serverless — CloudExpense"
prs.core_properties.subject="Apresentação académica · Turma Lecc31"
prs.core_properties.author="Ayrton Nhaguaga; Abelardo Miambo; Edmilson Conge"
BLACK=RGBColor(0,0,0)
WHITE=RGBColor(255,255,255)
LIGHT=RGBColor(246,246,246)
records=[]
def text(s,value,x,y,w,h,size=20,bold=False,align=None):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame
    tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01)
    tf.margin_top=tf.margin_bottom=Inches(.02)
    for i,row in enumerate(value.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=row;p.font.name="Arial";p.font.size=Pt(size)
        p.font.bold=bold;p.font.color.rgb=BLACK;p.space_after=Pt(8)
        if align is not None:p.alignment=align
    return box
def rect(s,x,y,w,h,fill=WHITE):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=fill
    sh.line.color.rgb=BLACK;sh.line.width=Pt(1)
    return sh
def line(s,x1,y1,x2,y2,width=1):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=BLACK;sh.line.width=Pt(width)
def arrow(s,x,y,w=.4):
    sh=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(.18))
    sh.fill.solid();sh.fill.fore_color.rgb=BLACK;sh.line.fill.background()
def pill(s,label,x,y,w):
    rect(s,x,y,w,.36)
    text(s,label,x+.08,y+.055,w-.16,.25,10,True)
def base(title,section,note,source=""):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=WHITE
    text(s,section.upper(),.62,.30,11,.24,10,True)
    text(s,title,.62,.85,12.06,.92,30,True)
    line(s,.64,1.86,12.69,1.86)
    line(s,.64,6.96,12.69,6.96,.6)
    text(s,"CLOUDEXPENSE  /  COMPUTAÇÃO SERVERLESS",.64,7.08,7,.17,8)
    text(s,f"{len(prs.slides):02d} / 24",11.5,7.05,1.17,.24,10,align=PP_ALIGN.RIGHT)
    if source:text(s,source,.66,6.64,11.97,.24,10)
    s.notes_slide.notes_text_frame.text=note
    records.append(dict(slide=len(prs.slides),title=title,section=section,notes=note,source=source))
    return s
def bullets(s,items,x,y,w,size=21,gap=.87):
    for i,item in enumerate(items):
        text(s,"•",x,y+i*gap,.22,.4,size,True)
        text(s,item,x+.32,y+i*gap,w-.32,gap-.08,size)
def card(s,n,title,body,x,y,w=3.8,h=2.7):
    rect(s,x,y,w,h)
    text(s,n,x+.22,y+.2,.8,.45,27,True)
    text(s,title,x+.22,y+.85,w-.44,.7,22,True)
    text(s,body,x+.22,y+1.6,w-.44,h-1.7,17)
def box(s,title,body,x,y,w,h=1.55):
    rect(s,x,y,w,h)
    text(s,title,x+.18,y+.17,w-.36,.70,18,True,PP_ALIGN.CENTER)
    text(s,body,x+.16,y+.80,w-.32,h-.85,15,align=PP_ALIGN.CENTER)
def picture(s,name,x,y,w,h,crop=None):
    path=ROOT/"imagens"/name
    iw,ih=Image.open(path).size
    if crop:
        pic=s.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h))
        pic.crop_left,pic.crop_top,pic.crop_right,pic.crop_bottom=crop
    else:
        scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
        pic=s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),Inches(ww),Inches(hh))
    return pic
def table(s,headers,rows,widths,x=.65,y=2.15,rowh=.76,size=17):
    total=sum(widths);rect(s,x,y,total,rowh,LIGHT)
    for r,row in enumerate([headers]+rows):
        xx=x
        for c,entry in enumerate(row):
            text(s,entry,xx+.13,y+r*rowh+.16,widths[c]-.26,rowh-.19,size,r==0)
            xx+=widths[c]
        if r:line(s,x,y+(r+1)*rowh,x+total,y+(r+1)*rowh,.6)
def cloud(s,x,y,w,h):
    sh=s.shapes.add_shape(MSO_SHAPE.CLOUD,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=WHITE;sh.line.color.rgb=BLACK;sh.line.width=Pt(1.6)

s=base("Computação Serverless","Apresentação académica",
"Apresentar os autores e o tema. O CloudExpense é o estudo de caso desenvolvido para demonstrar funções activadas por HTTP, persistência num serviço gerido e uma interface web moderna. O trabalho está funcional e validado localmente; a publicação remota ainda não foi realizada. Não indicar uma instituição ou curso que não tenham sido fornecidos.")
text(s,"CloudExpense",.65,2.15,8,1,50,True)
text(s,"Smart Personal Finance",.68,3.23,7,.50,25)
text(s,"Uma aplicação de gestão financeira pessoal\ncom Cloudflare Workers e Cloudflare D1.",.68,3.97,7.65,1.02,22)
text(s,"Ayrton Nhaguaga  ·  Abelardo Miambo  ·  Edmilson Conge",.68,5.55,11.9,.43,19,True)
text(s,"Turma: Lecc31     |     Docente: Sergio Sambo",.68,6.12,11.9,.35,17)
cloud(s,9.02,2.22,3.05,2.03)
text(s,"SERVERLESS",9.47,3.04,2.08,.40,17,True,PP_ALIGN.CENTER)
line(s,10.56,4.23,10.56,4.81)
box(s,"Worker + D1","Execução e persistência",9.02,4.83,3.05,1.45)

s=base("Percurso da apresentação","Roteiro",
"Organizar a exposição em quatro momentos. Primeiro definir o problema e os objectivos; depois explicar os fundamentos; em seguida mostrar a arquitectura e a aplicação; terminar com resultados, limitações e conclusão. Duração sugerida: 18 a 22 minutos, incluindo uma demonstração breve.")
for n,t,b,x,y in [
("01","Contexto e objectivos","Introdução, problema e justificativa.",.65,2.2),
("02","Fundamentos","Serverless, FaaS e responsabilidades.",6.84,2.2),
("03","Implementação","Arquitectura, dados, interface e segurança.",.65,4.26),
("04","Avaliação","Demonstração, testes e conclusão.",6.84,4.26)]:
    rect(s,x,y,5.84,1.78)
    text(s,n,x+.2,y+.2,.65,.5,28,True)
    text(s,t,x+1.02,y+.23,4.5,.4,23,True)
    text(s,b,x+1.02,y+.87,4.5,.65,18)

s=base("Introdução: do registo à informação","Contexto",
"Explicar que o projecto liga um problema de organização financeira a uma questão tecnológica. Não se apresenta uma pesquisa estatística sobre hábitos dos utilizadores; trata-se do cenário que orientou o desenvolvimento. O foco académico é compreender como a aplicação processa cada acção sem manter um servidor de aplicação tradicional.",
"Fonte: descrição e objectivos do projecto CloudExpense.")
bullets(s,["Registar receitas e despesas permite organizar os movimentos financeiros.","Categorias e relatórios transformam registos em informação útil.","Uma aplicação web exige interface, processamento e persistência."],.72,2.22,7.22,21,1.14)
box(s,"Questão tecnológica","Como executar essa lógica numa infraestrutura gerida?",8.65,2.45,3.88,2.2)
text(s,"Caso de estudo",8.68,5.04,3.9,.4,16,True)
text(s,"CloudExpense: uma aplicação completa para demonstrar serverless.",8.68,5.38,3.87,1.10,21)

s=base("Problema e pergunta de investigação","Problematização",
"Apresentar dois níveis do problema: o funcional, organizar e consultar movimentos, e o de engenharia, implementar a aplicação sem administrar um servidor permanente. A pergunta orienta a demonstração prática; o trabalho não pretende provar que serverless é sempre mais barato ou mais rápido.")
text(s,"Necessidade funcional",.72,2.23,5.7,.45,22,True)
bullets(s,["Centralizar receitas, despesas e categorias.","Calcular saldos e consultar o histórico com consistência."],.72,2.96,5.64,20,1.08)
text(s,"Desafio de implementação",6.99,2.23,5.7,.45,22,True)
bullets(s,["Executar lógica sob demanda.","Persistir dados e validar entradas sem um backend Node.js permanente."],6.99,2.96,5.62,20,1.08)
rect(s,.72,5.46,11.9,.88,LIGHT)
text(s,"Como construir uma aplicação financeira funcional com uma arquitectura serverless?",.96,5.55,11.43,.73,20,True)

s=base("Justificativa e relevância","Motivação",
"O projecto tem relevância prática, pedagógica e técnica. A gestão financeira oferece operações fáceis de observar, enquanto o percurso HTTP até ao banco permite explicar eventos, função e serviço gerido. Segurança e testes fazem parte da demonstração.")
card(s,"01","Aplicação prática","Um domínio claro: receitas, despesas, saldo e categorias.",.66,2.37,3.84,3.4)
card(s,"02","Aprendizagem","Tornar visível o fluxo entre interface, função e persistência.",4.76,2.37,3.84,3.4)
card(s,"03","Engenharia","Integrar validação, integridade, testes e documentação.",8.86,2.37,3.84,3.4)
text(s,"A contribuição é uma implementação demonstrável, acompanhada de evidências locais.",.78,6.08,11.8,.46,19)

s=base("Objectivo geral e objectivos específicos","Objectivos",
"Ler o objectivo geral e resumir os específicos em três grupos: implementar a API e persistência; construir a interface e relatórios; validar, observar e preparar a publicação. Todos se relacionam com os critérios de demonstração do trabalho.")
text(s,"Desenvolver uma aplicação web de gestão financeira pessoal baseada em serverless, com execução sob demanda e persistência de dados.",.73,2.16,11.91,1.40,25,True)
line(s,.74,3.71,12.57,3.71)
for n,t,b,x in [
("01","Implementar","API REST, regras de negócio e D1.",.73),
("02","Integrar","React, CRUD, filtros e gráficos reais.",4.83),
("03","Validar","Testes, logs e documentação de deploy.",8.93)]:
    text(s,n,x,4.04,.8,.62,32,True)
    text(s,t,x,4.91,3.69,.47,22,True)
    text(s,b,x,5.61,3.64,.72,20)

s=base("Metodologia de desenvolvimento","Método",
"Caracterizar o trabalho como desenvolvimento aplicado de um protótipo funcional, com verificação incremental. Os requisitos orientaram a construção; os testes verificaram regras e fluxos. Não houve estudo com participantes nem comparação experimental de desempenho entre plataformas.")
stages=[("01","Requisitos","Âmbito e regras"),("02","Modelação","Camadas e dados"),("03","Construção","API e interface"),("04","Validação","Testes e browser")]
for i,(n,t,b) in enumerate(stages):
    x=.72+i*3.12
    text(s,n,x,2.53,1,.68,34,True)
    box(s,t,b,x,3.4,2.64,1.57)
    if i<3:arrow(s,x+2.74,4.03,.27)
text(s,"Iteração: implementar → verificar → corrigir → verificar novamente",.79,5.60,11.7,.66,24,True)
text(s,"Artefactos: código, migrações, testes, documentação e apresentação.",.79,6.19,11.7,.3,16)

s=base("O que é Computação Serverless?","Fundamentos",
"Serverless significa que a infraestrutura de execução é gerida pelo fornecedor. Os servidores físicos continuam a existir. FaaS é a execução de funções em resposta a eventos; neste projecto, o evento é um pedido HTTP. O estado financeiro fica no D1, não na memória temporária do Worker. Não assumir que cada pedido cria uma máquina ou um processo novo.",
"Referências: [1] funcionamento de Workers; [2] visão geral do D1.")
text(s,"Infraestrutura gerida pelo fornecedor;\nlógica da aplicação activada por eventos.",.73,2.18,11.9,1.25,28,True)
box(s,"Evento","Pedido HTTP",.76,3.79,3.14,1.58)
arrow(s,4.12,4.49,.5)
box(s,"Função","Lógica no Worker",4.87,3.79,3.58,1.58)
arrow(s,8.68,4.49,.5)
box(s,"Persistência","Dados no D1",9.44,3.79,3.14,1.58)
text(s,"FaaS: Function as a Service",.79,5.72,6,.4,20,True)
text(s,"Serverless não significa ausência física de servidores.",.79,6.17,11.72,.35,18)

s=base("Hospedagem tradicional e serverless","Fundamentos",
"Comparar responsabilidades, evitando generalizações. Na hospedagem tradicional aqui considerada existe um processo de aplicação numa máquina provisionada. O modelo gerido retira tarefas de operação de infraestrutura, mas não elimina o trabalho com código, segurança, dados e custos. Ambos podem ser válidos conforme os requisitos.",
"Referência: [1]. Comparação conceptual; sem benchmark de custos ou desempenho.")
table(s,["Aspecto","Servidor tradicional","CloudExpense serverless"],[
["Execução","Processo numa máquina","Handler HTTP no Worker"],
["Infraestrutura","Provisionamento e manutenção","Execução gerida pela Cloudflare"],
["Escala","Instâncias e balanceamento","Gerida, sujeita a limites"],
["Dados","Banco local ou serviço externo","D1 gerido e persistente"],
["Programador","Código + operação do ambiente","Código, dados e segurança"],
],[2.18,4.72,5.12],rowh=.66,size=17)
text(s,"A responsabilidade muda de distribuição; não desaparece.",.8,6.29,11.65,.33,18,True)

s=base("Tecnologias e responsabilidades","Solução proposta",
"Explicar a stack pelo papel de cada conjunto, sem transformar a apresentação numa lista de bibliotecas. React trata da experiência; Hono organiza a API; D1 persiste. Wrangler executa e publica; testes validam. Node.js é usado para ferramentas de desenvolvimento, não como servidor permanente de produção.")
for x,y,title,body in [
(.72,2.2,"Interface","React + TypeScript + Vite\nTailwind CSS, Recharts e TanStack Query"),
(6.87,2.2,"Lógica da aplicação","Cloudflare Workers + Hono\nZod e API REST em JSON"),
(.72,4.2,"Persistência","Cloudflare D1 / SQLite\nSQL parametrizado e migrações"),
(6.87,4.2,"Desenvolvimento e testes","Wrangler, Vitest e Testing Library\nPlaywright para os fluxos no browser")]:
    rect(s,x,y,5.73,1.71)
    text(s,title,x+.22,y+.19,5.2,.44,23,True)
    text(s,body,x+.22,y+.82,5.2,.72,17)
text(s,"Node.js executa as ferramentas; o backend de produção executa no Worker.",.79,6.21,11.8,.36,17)

s=base("Arquitectura da aplicação","Arquitectura",
"Seguir as setas da esquerda para a direita. React comunica por HTTP REST com Hono, que chama services e repositories. O repository usa o binding DB. A resposta regressa em JSON. Em produção, Access protege o acesso e Static Assets serve React. A configuração remota está preparada, mas ainda não publicada.",
"Fonte: docs/ARCHITECTURE.md. Protecção Access e publicação remota ainda pendentes.")
nodes=[("Utilizador","Acção no browser",.67),("React","Interface",3.18),("Worker / Hono","API REST",5.69),("Service","Regras de negócio",8.20),("Repository","SQL parametrizado",10.71)]
for i,(t,b,x) in enumerate(nodes):
    box(s,t,b,x,2.17,1.96,1.60)
    if i<4:arrow(s,x+2.04,2.90,.34)
line(s,11.68,3.77,11.68,4.30)
sh=s.shapes.add_shape(MSO_SHAPE.CAN,Inches(10.71),Inches(4.3),Inches(1.96),Inches(1.42))
sh.fill.solid();sh.fill.fore_color.rgb=WHITE;sh.line.color.rgb=BLACK
text(s,"D1",11.07,4.82,1.2,.5,28,True,PP_ALIGN.CENTER)
line(s,10.69,5.03,4.20,5.03)
text(s,"Resposta JSON → interface actualizada",4.17,5.36,6.1,.61,20)
pill(s,"HTTP REST",3.62,4.14,1.75)
text(s,"Produção: Cloudflare Access + Workers Static Assets + API no mesmo domínio.",.81,6.16,11.86,.37,17)

s=base("O percurso de uma transacção","Execução sob demanda",
"Usar este slide para responder à pergunta 'onde está a função?'. O evento entra no handler do Worker, Hono selecciona a rota e a lógica é executada. A escrita é aguardada antes da confirmação. A interface invalida as consultas para obter os totais do backend. Um ambiente de execução pode ser reutilizado; os dados não dependem da sua memória.")
steps=[("01","Acção","O utilizador preenche e submete o formulário."),
("02","Requisição","React envia POST /api/transactions."),
("03","Validação","Worker verifica dados, montante e categoria."),
("04","Persistência","Repository executa INSERT parametrizado no D1."),
("05","Resposta","HTTP 201; a interface consulta novamente os totais.")]
for i,(n,t,b) in enumerate(steps):
    y=2.15+i*.81
    text(s,n,.80,y,.62,.44,24,True)
    text(s,t,1.73,y,2.43,.44,21,True)
    text(s,b,4.31,y,8.1,.62,20)
    if i<4:line(s,.83,y+.64,12.46,y+.64,.5)

s=base("Modelo de dados e integridade","Persistência",
"Mostrar a relação de uma categoria para muitas transacções. O esquema tem timestamps, índices e restrições adicionais; o diagrama mostra os campos essenciais. Explicar a conversão monetária sem arredondamento e a proibição de eliminar categorias utilizadas. O saldo é calculado, não duplicado numa coluna.",
"Fonte: worker/migrations/0001_initial.sql e shared/schemas.ts.")
rect(s,.80,2.22,4.52,2.68)
text(s,"CATEGORIES",1.01,2.44,4.1,.39,22,True)
text(s,"id  ·  chave primária\nname  ·  type  ·  color\ncreated_at  ·  updated_at",1.01,3.09,4.06,1.45,20)
rect(s,7.11,2.22,5.36,2.68)
text(s,"TRANSACTIONS",7.34,2.44,4.9,.39,22,True)
text(s,"id  ·  category_id (FK)\ntype  ·  amount  ·  description\ntransaction_date  ·  timestamps",7.34,3.09,4.86,1.45,20)
line(s,5.33,3.50,7.10,3.50,1.3)
text(s,"1",5.44,3.07,.3,.36,18,True)
text(s,"N",6.65,3.07,.3,.36,18,True)
text(s,"1.500,50 MT  →  150050 centavos (INTEGER)",.82,5.31,11.57,.53,27,True)
text(s,"Chaves estrangeiras + CHECK + triggers + consultas parametrizadas.",.84,6.05,11.65,.45,20)

s=base("Dashboard: informação financeira visível","Interface",
"Mostrar a captura real obtida num teste local. Os valores são registos temporários de teste, removidos depois da validação. Destacar cartões, evolução por mês e despesas por categoria. Todos recebem valores da API. O dashboard mostra o ano seleccionado; o saldo também se refere a esse período.",
"Imagem: captura real do CloudExpense em teste local; dados de demonstração, não dados pessoais.")
picture(s,"dashboard-desktop.png",.71,2.05,9.02,4.43,crop=(0,.13,0,.34))
rect(s,.70,2.04,9.04,4.45).fill.background()
text(s,"Dados reais\nda API",10.09,2.39,2.5,1.04,24,True)
bullets(s,["Saldo do ano","Receitas e despesas","Distribuição por categoria"],10.08,3.88,2.55,17,.74)

s=base("Funcionalidades e experiência de utilização","Interface",
"A aplicação permite criar, ler, editar e eliminar movimentos, e gerir categorias com restrições. Os filtros combinam-se com paginação. Mostrar a captura mobile e explicar que as tabelas largas têm scroll interno. Não afirmar uma auditoria completa de acessibilidade: foram implementados foco, labels, teclado e feedback, e validados fluxos no browser.",
"Imagem: captura mobile do teste E2E. Fonte funcional: README.md e docs/TESTING.md.")
picture(s,"dashboard-mobile.png",.90,2.13,3.09,4.26,crop=(0,0,0,.72))
rect(s,.90,2.13,3.09,4.26).fill.background()
bullets(s,[
"Receitas e despesas: criar, editar e eliminar com confirmação.",
"Categorias: nome, tipo e cor, preservando relações existentes.",
"Pesquisa, filtros combinados e paginação.",
"Layout responsivo, validação, loading e mensagens de erro."
],4.51,2.26,7.87,22,1.00)

s=base("Segurança em várias camadas","Segurança",
"Distinguir protecção implementada de configuração externa. Zod valida em ambos os lados, mas o backend é a autoridade. SQL parametrizado e constraints protegem dados. O middleware de produção valida assinatura e claims JWT. Ainda é preciso criar a política Access real na conta. O MVP é um espaço partilhado, sem contas individuais.",
"Referência: [3] validação de JWT Access. Evidência local: testes do Worker.")
for n,t,b,x in [
("01","Entrada","Zod, limites de corpo e rejeição de origens de escrita indevidas.",.72),
("02","Dados","SQL parametrizado, chaves estrangeiras e regras de integridade.",4.83),
("03","Acesso","JWT Access validado; produção bloqueia se faltar configuração.",8.94)]:
    card(s,n,t,b,x,2.23,3.68,3.60)
text(s,"MVP académico partilhado: não destinado a dados financeiros privados reais.",.81,6.12,11.69,.40,18,True)

s=base("Demonstração: do movimento ao saldo","Validação funcional",
"Executar, se houver tempo, os passos no browser. Primeiro criar receita de 25000 MT na categoria Salário; depois despesa de 5000 MT em Alimentação, com datas no mesmo ano seleccionado. Confirmar os cartões e recarregar. Os testes E2E também fizeram edição, eliminação e gestão de categorias.",
"Fonte: testes locais de integração e E2E. Montantes temporários de demonstração.")
for title,value,x in [("Receita","25.000 MT",.78),("Despesa","5.000 MT",5.0),("Saldo","20.000 MT",9.21)]:
    rect(s,x,2.39,3.33,1.81)
    text(s,title,x+.16,2.65,3.01,.43,20,True,PP_ALIGN.CENTER)
    text(s,value,x+.12,3.32,3.09,.53,29,True,PP_ALIGN.CENTER)
text(s,"−",4.33,3.12,.42,.57,34,True)
text(s,"=",8.55,3.12,.42,.57,34,True)
bullets(s,["Recarregar a página: os dados permanecem no D1 local.","Editar ou eliminar: os totais são recalculados pelo backend."],.89,4.86,11.46,22,.81)

s=base("Testes e resultados obtidos","Avaliação",
"Os números são do relatório de testes registado em 11 de Outubro de 2026. São 55 testes unitários/de integração e 2 percursos end-to-end, totalizando 57 casos executados. Distinguir testes de componente com mocks HTTP dos testes Worker/D1 no runtime local. Estes resultados não comprovam desempenho ou disponibilidade na cloud.",
"Fonte: docs/TESTING.md — execução local em 11/10/2026.")
for n,t,b,x in [
("23","Frontend","Formulários, moeda, estados e cliente HTTP.",.75),
("32","Worker + D1","CRUD, integridade, filtros, relatórios e Access.",4.88),
("02","Browser / E2E","Percursos completos em desktop e mobile.",9.01)]:
    rect(s,x,2.28,3.59,2.90)
    text(s,n,x+.22,2.45,3.15,1.08,58,True)
    text(s,t,x+.23,3.59,3.14,.43,22,True)
    text(s,b,x+.23,4.26,3.08,.74,17)
text(s,"Também aprovados: TypeScript, migração local, build e Wrangler --dry-run.",.80,5.70,11.71,.6,21,True)
text(s,"Limite da evidência: sem teste de carga nem validação de um deploy remoto.",.81,6.25,11.75,.32,16)

s=base("Como observar a execução serverless","Observabilidade",
"Mostrar a aba Network durante a demonstração e relacionar o pedido com o log. O requestId permite correlacionar eventos. Não mostrar uma latência inventada no slide: o painel mede-a ao vivo. A duração HTTP inclui rede e processamento e não corresponde ao tempo de CPU facturado. Métricas remotas só estarão disponíveis depois de publicar.")
text(s,"No ambiente local",.81,2.25,5.55,.5,24,True)
bullets(s,["Network: pedido HTTP e resposta JSON.","Health check: Worker consulta o D1.","Logs: método, rota, status e requestId."],.82,3.00,5.56,20,.92)
text(s,"Depois da publicação",7.04,2.25,5.46,.5,24,True)
bullets(s,["Wrangler tail e Workers Observability.","Invocações, erros e consumo na conta.","Verificar Access com utilizadores autorizados."],7.04,3.00,5.44,20,.92)
rect(s,.82,6.03,11.72,.51,LIGHT)
text(s,"Tempo HTTP no cliente ≠ tempo de CPU do Worker.",1.01,6.13,11.32,.30,18,True)

s=base("Vantagens, custos e limitações","Discussão",
"Relacionar vantagens com escolhas observadas, sem prometer poupança medida. O modelo de cobrança depende do plano; não inserir preços fixos. Neste projecto todos os pedidos passam pelo Worker para validação de Access, incluindo assets. D1 tem limites próprios; escalabilidade do compute não elimina gargalos nos dados.",
"Referências: [4] preços Workers; [5] preços D1. Não foi realizada análise comparativa de custos.")
text(s,"Vantagens",.80,2.24,5.53,.50,24,True)
bullets(s,["Menos operação directa de servidores.","Execução gerida e integração com D1.","Desenvolvimento local antes de publicar."],.80,3.0,5.47,20,.91)
text(s,"Limitações e cuidados",6.98,2.24,5.60,.50,24,True)
bullets(s,["Quotas, limites e dependência do fornecedor.","Custos variam com utilização e plano.","Segurança, dados e testes continuam a cargo da equipa."],6.98,3.0,5.55,20,.91)
text(s,"Workers: pedidos e CPU.  D1: leituras, escritas e armazenamento.",.82,6.14,11.78,.38,18)

s=base("Estado actual e caminho para publicação","Execução e deploy",
"Ser preciso sobre o estado do trabalho: funciona no computador com Wrangler e D1 local. O dry-run apenas valida o empacotamento. Para publicar, falta autenticar na conta, criar a base remota, preencher o ID, configurar domínio e Access, aplicar migrações e publicar com autorização. Dados locais não são transferidos automaticamente.",
"Fonte: docs/DEPLOYMENT.md. Nenhuma publicação remota foi realizada.")
pill(s,"CONCLUÍDO LOCALMENTE",.79,2.22,3.18)
bullets(s,["Frontend, API e persistência D1.","Testes, build e pacote de publicação.","Worker com Static Assets no mesmo servidor local."],.8,3.00,5.35,20,.96)
pill(s,"CONFIGURAÇÃO EXTERNA",6.95,2.22,3.28)
bullets(s,["Criar D1 remoto e configurar o ID real.","Configurar domínio e Cloudflare Access.","Aplicar migrações e publicar com Wrangler."],6.95,3.00,5.59,20,.96)
text(s,"Local: npm.cmd run dev  →  http://127.0.0.1:5173",.82,6.17,11.79,.38,19,True)

s=base("Conclusão e trabalhos futuros","Conclusão",
"Retomar o objectivo geral. A implementação demonstra execução orientada a pedidos, regras no Worker e persistência independente no D1 local. Os testes sustentam o funcionamento nos fluxos verificados. A demonstração remota exige a configuração e publicação ainda pendentes. Melhorias futuras não devem ser apresentadas como funcionalidades já existentes.")
text(s,"O CloudExpense materializa os conceitos de serverless numa aplicação funcional e verificável localmente.",.77,2.21,11.8,1.13,29,True)
line(s,.82,3.70,12.48,3.70)
text(s,"Contribuição",.82,4.08,5.5,.48,23,True)
text(s,"Interface, função e persistência integradas, com regras de negócio, segurança e evidências de teste.",.82,4.84,5.51,1.4,21)
text(s,"Próximos passos",6.98,4.08,5.5,.48,23,True)
text(s,"Publicar e medir o ambiente remoto.\nAvaliar acessibilidade e carga.\nEstudar contas e isolamento de dados.",6.98,4.84,5.55,1.5,21)

s=base("Referências e fontes das evidências","Referências",
"As referências técnicas são documentação oficial Cloudflare, consultada em 11 de Outubro de 2026. Os números de testes e capturas vêm do projecto local. Os links completos estão nas notas e no guião. Os diagramas foram elaborados para esta apresentação.",
"Consulta: 11/10/2026. Capturas e diagramas: elaboração própria a partir do projecto.")
refs=[
("[1] Cloudflare — How Workers works","https://developers.cloudflare.com/workers/reference/how-workers-works/"),
("[2] Cloudflare — D1 overview","https://developers.cloudflare.com/d1/"),
("[3] Cloudflare One — Validate JWTs","https://developers.cloudflare.com/cloudflare-one/access-controls/applications/http-apps/authorization-cookie/validating-json/"),
("[4] Cloudflare — Workers pricing","https://developers.cloudflare.com/workers/platform/pricing/"),
("[5] Cloudflare — D1 pricing","https://developers.cloudflare.com/d1/platform/pricing/"),
("[6] CloudExpense — documentação e testes","docs/ARCHITECTURE.md · SERVERLESS.md · TESTING.md")]
for i,(label,url) in enumerate(refs):
    y=2.16+i*.67
    shape=text(s,label,.81,y,11.74,.37,20,True)
    if url.startswith("https"):
        for p in shape.text_frame.paragraphs:
            for run in p.runs:run.hyperlink.address=url
    short=url if i in [1,3,4,5] else ("developers.cloudflare.com/workers/reference/how-workers-works/" if i==0 else "developers.cloudflare.com · Cloudflare One · Access · Validate JWTs")
    text(s,short,.84,y+.37,11.73,.25,12)
s.notes_slide.notes_text_frame.text += "\n\n" + "\n".join(f"{a}\n{b}" for a,b in refs)

s=base("Questões e discussão","Encerramento",
"Agradecer e convidar perguntas. Preparar respostas sobre a diferença entre FaaS e serverless, localização da persistência, validação monetária, ausência de contas, limites do teste local e configuração de Access. Se solicitado, abrir a aplicação ou a documentação do repositório.")
text(s,"A infraestrutura é gerida.\nA responsabilidade pelo software continua.",.79,2.28,11.86,1.70,36,True)
text(s,"Obrigado.",.81,4.55,11.7,.73,38,True)
text(s,"Ayrton Nhaguaga  ·  Abelardo Miambo  ·  Edmilson Conge",.82,5.72,11.72,.4,19)
text(s,"Lecc31   |   Docente: Sergio Sambo",.82,6.23,11.74,.3,16)

assert len(prs.slides)==24
out=DOCS/"CloudExpense_Computacao_Serverless.pptx"
# Keep PowerPoint hyperlink colours black, including visited links.
from lxml import etree
ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
for part in prs.part.package.iter_parts():
    if str(part.partname).startswith("/ppt/theme/") and str(part.partname).endswith(".xml"):
        theme = etree.fromstring(part.blob)
        for tag in ("hlink", "folHlink"):
            node = theme.find(".//{" + ns + "}clrScheme/{" + ns + "}" + tag)
            if node is not None:
                for child in list(node): node.remove(child)
                etree.SubElement(node, "{" + ns + "}srgbClr", val="000000")
        part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)
prs.save(out)
issues=[]
for idx,slide in enumerate(prs.slides,1):
    for sh in slide.shapes:
        if sh.left < 0 or sh.top < 0 or sh.left+sh.width>prs.slide_width+10 or sh.top+sh.height>prs.slide_height+10:
            issues.append(f"Slide {idx}: shape outside canvas")
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.color.type is not None and r.font.color.rgb!=BLACK:
                        issues.append(f"Slide {idx}: text is not black")
if issues:raise RuntimeError("; ".join(issues))
guide=["# Guião da apresentação — CloudExpense","",
"Autores: Ayrton Nhaguaga, Abelardo Miambo e Edmilson Conge. Turma: Lecc31. Docente: Sergio Sambo.","",
"## Como usar","",
"- Abrir ../CloudExpense_Computacao_Serverless.pptx no PowerPoint e iniciar com F5.",
"- Formato 16:9, fundo branco e texto preto. Textos e diagramas editáveis; capturas nas cores originais da aplicação.",
"- Duração sugerida: 18–22 minutos, incluindo uma demonstração breve.",
"- Notas do apresentador incorporadas no PowerPoint.",
"- Os resultados apresentados são locais: não afirmar que o sistema já foi publicado.",
"- Não foram inventados instituição, curso ou dados de desempenho.","",
"## Divisão sugerida da exposição","",
"- Ayrton Nhaguaga: slides 1–8 — contexto, objectivos e fundamentos.",
"- Abelardo Miambo: slides 9–16 — tecnologias, arquitectura e implementação.",
"- Edmilson Conge: slides 17–24 — demonstração, resultados e conclusão.",
"- A divisão é uma sugestão e pode ser alterada.","",
"## Demonstração ao vivo","",
"1. Na raiz: npm.cmd run db:migrate e npm.cmd run dev.",
"2. Abrir http://127.0.0.1:5173 e seleccionar o ano das transacções.",
"3. Criar receita de 25000 MT e despesa de 5000 MT com categorias compatíveis.",
"4. Numa base sem outros movimentos desse ano, verificar saldo de 20000 MT.",
"5. Recarregar, editar, eliminar e observar as respostas na aba Network.",
"6. Mostrar os logs do Worker, sem inventar métricas de CPU.",
"7. Se já houver movimentos, explicar a diferença de saldo de 20000 MT.","",
"## Notas por slide",""]
for r in records:guide += [f"### {r['slide']:02d}. {r['title']}","",r["notes"],"",f"Fonte no slide: {r['source']}" if r["source"] else "",""]
guide+=["## Referências completas",""]+[f"- {a}: {b}" for a,b in refs]
guide+=["","## Regenerar","","Requer Python, python-pptx (1.0.2 utilizado) e Pillow.","",
"Comando: python docs/apresentacao/gerar_apresentacao.py","",
"O script lê imagens/ e grava o PPTX em docs/. Exportar PDF no PowerPoint para preservar a paginação.",""]
(ROOT/"GUIAO.md").write_text("\n".join(guide),encoding="utf-8")
(ROOT/"slides.json").write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding="utf-8")
print(f"Created {out.name}: {len(prs.slides)} slides; all shapes in canvas; native text black.")